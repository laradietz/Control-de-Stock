"""exportador.py - Exportación a Excel, PDF de inventario, boletas y reportes de caja."""
import os
import sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER, TA_RIGHT

if os.environ.get("FERRETERIA_BASE_DIR"):
    _base = os.environ["FERRETERIA_BASE_DIR"]
elif getattr(sys, 'frozen', False):
    _base = os.path.dirname(sys.executable)
else:
    _base = os.path.dirname(os.path.abspath(__file__))

EXPORT_DIR = os.path.join(_base, "exports")

def _dir():
    os.makedirs(EXPORT_DIR, exist_ok=True)

def exportar_excel(productos, nombre_archivo=None):
    _dir()
    if not nombre_archivo:
        nombre_archivo = f"inventario_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    ruta = os.path.join(EXPORT_DIR, nombre_archivo)
    wb = Workbook(); ws = wb.active; ws.title = "Inventario"
    encabezados = ["Código","Nombre","Categoría","Precio ($)","Cantidad","Stock Mínimo","Proveedor","Estado"]
    ws.append(encabezados)
    hf = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    for col in range(1, len(encabezados)+1):
        c = ws.cell(row=1, column=col)
        c.fill = hf; c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center")
    bf = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    for i, p in enumerate(productos, 2):
        estado = "STOCK BAJO" if p.stock_bajo() else "OK"
        ws.append([p.codigo,p.nombre,p.categoria_nombre,p.precio,p.cantidad,p.stock_minimo,p.proveedor_nombre,estado])
        if p.stock_bajo():
            for col in range(1,9): ws.cell(row=i,column=col).fill = bf
    for i,w in enumerate([12,35,20,12,10,12,25,12],1):
        ws.column_dimensions[chr(64+i)].width = w
    wb.save(ruta); return ruta

def exportar_pdf(productos, nombre_archivo=None, titulo="Reporte de Inventario - Ferretería Gian"):
    _dir()
    if not nombre_archivo:
        nombre_archivo = f"inventario_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    ruta = os.path.join(EXPORT_DIR, nombre_archivo)
    doc = SimpleDocTemplate(ruta, pagesize=landscape(A4),
                             leftMargin=1.5*cm, rightMargin=1.5*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
    estilos = getSampleStyleSheet(); elems = []
    elems.append(Paragraph(titulo, estilos["Title"]))
    elems.append(Paragraph(f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}", estilos["Normal"]))
    elems.append(Spacer(1, 0.4*cm))
    data = [["Código","Nombre","Categoría","Precio ($)","Cantidad","Stock Mín.","Proveedor","Estado"]]
    for p in productos:
        data.append([p.codigo,p.nombre,p.categoria_nombre,f"{p.precio:,.2f}",str(p.cantidad),str(p.stock_minimo),p.proveedor_nombre,"STOCK BAJO" if p.stock_bajo() else "OK"])
    t = Table(data, repeatRows=1)
    estilo = [("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2F5496")),
              ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
              ("FONTSIZE",(0,0),(-1,-1),8),("GRID",(0,0),(-1,-1),0.5,colors.grey),
              ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F2F2F2")])]
    for i,p in enumerate(productos,1):
        if p.stock_bajo(): estilo.append(("BACKGROUND",(0,i),(-1,i),colors.HexColor("#FFC7CE")))
    t.setStyle(TableStyle(estilo)); elems.append(t)
    doc.build(elems); return ruta

def generar_boleta_pdf(id_venta, detalles, total, fecha, cliente_nombre="Cliente General",
                       nombre_archivo=None, forma_pago="efectivo", descuento=0):
    _dir()
    if not nombre_archivo:
        nombre_archivo = f"boleta_{id_venta}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    ruta = os.path.join(EXPORT_DIR, nombre_archivo)
    doc = SimpleDocTemplate(ruta, pagesize=A4,
                             leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    estilos = getSampleStyleSheet()
    titulo_style = ParagraphStyle("tit", parent=estilos["Title"], fontSize=18, textColor=colors.HexColor("#2F5496"))
    sub_style = ParagraphStyle("sub", parent=estilos["Normal"], fontSize=10, textColor=colors.grey)
    normal = estilos["Normal"]
    elems = []
    elems.append(Paragraph("Ferretería Gian", titulo_style))
    elems.append(Paragraph("Sistema de Control de Stock", sub_style))
    elems.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#2F5496")))
    elems.append(Spacer(1, 0.4*cm))
    formas = {"efectivo":"Efectivo","tarjeta":"Tarjeta","transferencia":"Transferencia","cuenta_corriente":"Cuenta corriente (fiado)"}
    info_data = [
        [Paragraph(f"<b>Boleta N°:</b> {id_venta:05d}", normal),
         Paragraph(f"<b>Fecha:</b> {fecha}", normal)],
        [Paragraph(f"<b>Cliente:</b> {cliente_nombre}", normal),
         Paragraph(f"<b>Forma de pago:</b> {formas.get(forma_pago, forma_pago)}", normal)],
    ]
    info_t = Table(info_data, colWidths=[9*cm, 7*cm])
    info_t.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("GRID",(0,0),(-1,-1),0,colors.white)]))
    elems.append(info_t); elems.append(Spacer(1, 0.4*cm))
    data = [["Descripción","Precio unit.","Cant.","Subtotal"]]
    subtotal = sum(d.subtotal for d in detalles)
    for d in detalles:
        data.append([d.nombre_producto, f"${d.precio_unitario:,.2f}", str(d.cantidad), f"${d.subtotal:,.2f}"])
    if descuento > 0:
        data.append(["","",f"Descuento {descuento:.0f}%", f"-${subtotal * descuento/100:,.2f}"])
    data.append(["","","<b>TOTAL</b>", f"<b>${total:,.2f}</b>"])
    det_t = Table(data, colWidths=[9*cm, 2.8*cm, 2.2*cm, 2.8*cm], repeatRows=1)
    det_style = [
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2F5496")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),9),("GRID",(0,0),(-1,-2),0.5,colors.grey),
        ("ROWBACKGROUNDS",(0,1),(-1,-2),[colors.white,colors.HexColor("#F2F2F2")]),
        ("BACKGROUND",(0,-1),(-1,-1),colors.HexColor("#EAF1FB")),
        ("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold"),("FONTSIZE",(0,-1),(-1,-1),11),
        ("ALIGN",(1,0),(-1,-1),"RIGHT"),
    ]
    det_t.setStyle(TableStyle(det_style)); elems.append(det_t)
    elems.append(Spacer(1, 1*cm))
    elems.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey))
    elems.append(Paragraph("¡Gracias por su compra!", ParagraphStyle("grac",parent=estilos["Normal"],alignment=TA_CENTER,fontSize=10,textColor=colors.grey)))
    doc.build(elems); return ruta

def exportar_caja_pdf(resumen, nombre_archivo=None):
    _dir()
    if not nombre_archivo:
        nombre_archivo = f"caja_{resumen['fecha']}.pdf"
    ruta = os.path.join(EXPORT_DIR, nombre_archivo)
    doc = SimpleDocTemplate(ruta, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    estilos = getSampleStyleSheet(); elems = []
    elems.append(Paragraph("Ferretería Gian — Cierre de Caja", estilos["Title"]))
    elems.append(Paragraph(f"Fecha: {resumen['fecha']}", estilos["Normal"]))
    elems.append(Spacer(1, 0.5*cm))
    data = [["Concepto","Valor"],
            ["Ventas realizadas", str(resumen["cantidad_ventas"])],
            ["Ventas anuladas", str(resumen["anuladas"])],
            ["Total recaudado", f"${resumen['total']:,.2f}"]]
    t = Table(data, colWidths=[10*cm, 6*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2F5496")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),11),("GRID",(0,0),(-1,-1),0.5,colors.grey),
        ("BACKGROUND",(0,-1),(-1,-1),colors.HexColor("#EAF1FB")),
        ("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold"),
    ]))
    elems.append(t)
    if resumen["top_productos"]:
        elems.append(Spacer(1, 0.5*cm))
        elems.append(Paragraph("Top productos vendidos", estilos["Heading2"]))
        data2 = [["Producto","Unidades vendidas"]] + [[r[0], str(r[1])] for r in resumen["top_productos"]]
        t2 = Table(data2, colWidths=[12*cm, 4*cm])
        t2.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2F5496")),
                                  ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),10),
                                  ("GRID",(0,0),(-1,-1),0.5,colors.grey)]))
        elems.append(t2)
    doc.build(elems); return ruta
